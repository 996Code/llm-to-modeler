"""文档解析与结构感知切块 —— SDK 通用文档设施。

【模块定位】
从 knowledge_graph 插件下沉的纯文档能力:文件字节 → 纯文本 → 结构感知
切块。零领域知识——不认识知识库/图谱/向量,输出只有 seq/text/char_count
纯文本结构,任何需要"解析文档再处理"的插件(知识图谱/RAG/报表语料
入库…)都可复用。

【支持格式】md/txt(UTF-8/GBK/BOM/宽字符容错解码)、pdf(pypdf,无文字层
的扫描页自动 OCR 兜底)、docx(python-docx,标题层级还原为 # 前缀)。
驱动库延迟 import——不装 pypdf/python-docx 的环境仍可用 md/txt 能力;
OCR 兜底(pypdfium2+rapidocr)同为可选依赖,未安装时保持纯文字层行为。

【安全边界】
- 扩展名白名单(ALLOWED_EXTENSIONS)+ 解析失败抛 ValueError(调用方
  上传预检当场拒绝);
- _decode_text 对 UTF-16/乱码做探测拒收,不让垃圾文本流入下游;
- 出口统一剥离 C0 控制字符与孤立代理对:坏字体映射会把字形映到
  控制码点或只吐出 UTF-16 代理对的一半(JJF 京 3024-2023 实测:
  均值符号 X̄ 被提取成 X\x00,另有孤立 \\ud835)——前者 PG 拒收,
  后者无法 UTF-8 编码,进库/进 LLM/进检索全是祸患。
"""
import io
import logging
import re
import threading
import unicodedata
from typing import Any, Dict, List

logger = logging.getLogger(__name__)

# 扩展名白名单(上传校验与解析分发共用)
ALLOWED_EXTENSIONS = {".md", ".markdown", ".txt", ".pdf", ".docx"}

MIME_BY_EXT = {
    ".md": "text/markdown", ".markdown": "text/markdown",
    ".txt": "text/plain", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

# 分页符和垂直制表符保留为分隔空白,避免相邻正文粘连。
_CTRL_RE = re.compile(r"[\x00-\x08\x0e-\x1f\x7f]")
_TEXT_SEPARATORS = str.maketrans({"\x0b": "\n", "\x0c": "\n"})

# UTF-16 代理对区间:坏 CMap 常只吐出代理对的一半(孤立代理)——无法
# UTF-8 编码,psycopg 写入时直接抛 UnicodeEncodeError
_SURROGATE_RE = re.compile(r"[\ud800-\udfff]")

# 文字层垃圾占比超过该值的页视为"文字层不可信",交 OCR 从图像重识别;
# 低于它只剥离坏字符保留文字层——整页重 OCR 等于拿 99% 的好文本去赌
# 识别质量(实测坏字体映射页密度仅 0.4%,全文档 5700 字符只有 4 个)
_CTRL_GARBAGE_RATIO = 0.02

# 短文字层页的扫描件判定:单张图片对象覆盖版面比例达到该值才算
# "整页图"(扫描页/位图图表页);向量图与普通排版页不含这种对象
_SCAN_IMAGE_COVERAGE = 0.6

_HEADING_RE = re.compile(r"^(#{1,6})\s+")

# 纯标题缓冲(目录区)被正文到达触发结算的最小长度:低于此值的标题串
# (如单独一行"目录")不值得独立成块,并入后续正文
_HEADING_FLUSH_MIN = 40

# 通用章节标题(纯文本书籍的软边界):"第一章 xxx"/"第 12 节"/"序章"/"尾声"等。
# 约束:标题行须短(<40 字)且行首允许全角空格缩进——正文里"第二天"这类
# 普通叙述不会被误判(不带"章/节/卷/部"量词单位)。
_CHAPTER_RE = re.compile(
    r"^[\s\u3000]*(第[一二三四五六七八九十百千零〇\d]{1,7}[章節节卷部回])"
    r"([\s\u3000]+\S[^。!?!?\n]{0,29})?$"
)
_SPECIAL_CHAPTER_RE = re.compile(
    r"^[\s\u3000]*(序章|序言|楔子|前言|引子|尾声|後記|后记|番外|主目录|目录|正文|附录)"
    r"([\s\u3000:.:]\S.{0,30})?$")


def _is_heading_line(stripped: str) -> bool:
    """是否为结构标题(markdown # 或通用章节标题)——切块软边界判定。"""
    if _HEADING_RE.match(stripped):
        return True
    if len(stripped) > 40:
        return False
    return bool(_CHAPTER_RE.match(stripped) or _SPECIAL_CHAPTER_RE.match(stripped))


def allowed_extension(filename: str) -> bool:
    name = (filename or "").lower()
    return any(name.endswith(ext) for ext in ALLOWED_EXTENSIONS)


def mime_for(filename: str) -> str:
    name = (filename or "").lower()
    for ext, mime in MIME_BY_EXT.items():
        if name.endswith(ext):
            return mime
    return "application/octet-stream"


# ── 解析:文件字节 → 纯文本 ─────────────────────────────────

def parse_to_text(filename: str, data: bytes, *, ocr: bool = True,
                  stats: Dict[str, int] = None,
                  on_ocr_progress=None,
                  ocr_engine: str = "rapidocr",
                  page_ocr=None) -> str:
    """按扩展名解析文件为纯文本(无法解析抛 ValueError,由任务层记失败)。

    Args:
        ocr: PDF 无文字层页是否走 OCR 兜底。上传预检传 False——预检只
            验坏文件,OCR 秒级耗时留给导入任务(结果不丢弃,也不阻塞上传)。
        stats: 可选出参字典,调用方传入后填充 pages/ocrPages/ocrFilled/
            ocrEngine/ctrlStripped(任务日志据此告知用户扫描页识别情况、
            所用引擎与垃圾字符清洗量)。
        on_ocr_progress: 可选回调 (done, total),OCR 每完成一页调用一次
            (任务层用它刷新进度,让用户看到 OCR 在动而不是"卡死")。
        ocr_engine: 扫描页识别引擎。rapidocr(默认,本地 PP-OCR)/llm
            (调用方注入 page_ocr 回调接多模态大模型——SDK 零领域知识,
            不认识 LLM 客户端,只认"图片进文本出"的函数)。
        page_ocr: engine=llm 时的识别回调 (PIL.Image) -> str。

    出口统一 NFKC 归一化(P2#12 坑位②):PDF 提取常带兼容字符
    (民法典实测出现康熙部首 ⼀/⼱ 而非规范字 一/巾),下游切块/
    实体抽取/全文检索按字面匹配,不归一则同一字两种形态永不相等。
    """
    name = (filename or "").lower()
    if name.endswith((".md", ".markdown", ".txt")):
        return _nfkc(_decode_text(data))
    if name.endswith(".pdf"):
        return _nfkc(_parse_pdf(data, ocr=ocr, stats=stats,
                               on_ocr_progress=on_ocr_progress,
                               ocr_engine=ocr_engine, page_ocr=page_ocr))
    if name.endswith(".docx"):
        return _nfkc(_parse_docx(data))
    raise ValueError(f"不支持的文件类型: {filename}(支持 md/txt/pdf/docx)")


def _nfkc(text: str) -> str:
    """出口统一归一化:代理对修复 → NFKC → C0 剥离。

    不变量:出口文本可安全 UTF-8 编码且不含 C0 控制符。代理对修复放在
    NFKC 前——配对合并出的星面字符(数学斜体 𝑋)随后被 NFKC 折叠为
    ASCII 等价(X),下游检索按字面可命中。
    """
    if _SURROGATE_RE.search(text):
        text = _fix_surrogates(text)[0]
    return _CTRL_RE.sub("", unicodedata.normalize("NFKC", text)).translate(_TEXT_SEPARATORS)


def _fix_surrogates(text: str):
    """修复坏 CMap 吐出的 UTF-16 代理对:相邻高低配对合并为真实字符,
    孤立的(另一半丢失)剔除。

    Returns: (修复后文本, 剔除的孤立代理数)。
    """
    out = []
    dropped = 0
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ("\ud800" <= ch <= "\udbff" and i + 1 < n
                and "\udc00" <= text[i + 1] <= "\udfff"):
            out.append(chr(0x10000 + ((ord(ch) - 0xD800) << 10)
                           + (ord(text[i + 1]) - 0xDC00)))
            i += 2
        elif "\ud800" <= ch <= "\udfff":
            dropped += 1
            i += 1
        else:
            out.append(ch)
            i += 1
    return "".join(out), dropped


def _decode_text(data: bytes) -> str:
    """UTF-8 优先,失败回退 GBK(中文办公环境常见),再失败用 errors=replace。

    BOM/零字节探测:UTF-16/UTF-32 字节流常能"成功"通过 GBK 解码成含 NUL
    的乱码(不抛异常),这类垃圾进切块再进 LLM 就是一张废图谱——所以
    GBK 解出来后若含 NUL 或高占比替换符,按解码失败处理。
    """
    if data[:3] == b"\xef\xbb\xbf":
        return data[3:].decode("utf-8", errors="replace")
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    if data[:4] in (b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff"):
        return data.decode("utf-32", errors="replace")
    if b"\x00" in data[:200]:
        # 无 BOM 但含零字节:未知宽字符编码,按 UTF-16 兜底再校验
        guess = data.decode("utf-16", errors="ignore")
        if guess and "\x00" not in guess:
            return guess
        raise ValueError("无法识别的文本编码(疑似宽字符/二进制内容)")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = data.decode("gbk")
            if "\x00" in text:
                raise ValueError("无法识别的文本编码(疑似宽字符/二进制内容)")
            return text
        except UnicodeDecodeError:
            fallback = data.decode("utf-8", errors="replace")
            # 替换符占比过高 = 实际上没解出来,别把乱码当文本
            if fallback.count("\ufffd") > len(fallback) * 0.05:
                raise ValueError("无法识别的文本编码")
            return fallback


def _parse_pdf(data: bytes, ocr: bool = True, stats: Dict[str, int] = None,
               on_ocr_progress=None, ocr_engine: str = "rapidocr",
               page_ocr=None) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    texts = []
    ctrl_stripped = 0        # 全文档剥离的垃圾字符数(stats 留痕给任务日志)
    garbage: List[int] = []  # 文字层垃圾占比高的页 → OCR 从图像重识别
    for i, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:  # 单页损坏不整体失败
            logger.warning(f"pdf 第 {i + 1} 页解析失败(跳过): {e}")
            text = ""
        n_ctrl = len(_CTRL_RE.findall(text))
        if n_ctrl:
            ctrl_stripped += n_ctrl
            if n_ctrl / max(len(text), 1) > _CTRL_GARBAGE_RATIO:
                garbage.append(i)
            text = _CTRL_RE.sub("", text)
        if _SURROGATE_RE.search(text):
            text, dropped = _fix_surrogates(text)
            ctrl_stripped += dropped
        texts.append(text)
    # OCR 候选:①无/近无文字层的页——空页直接候选(无可失去);非空短页
    # 先结构复核(_filter_scan_pages),纯短文本页(章末空白/整页图表的
    # 短图注)文字层可信,不再误送 OCR;②文字层垃圾页(并集去重)
    sparse = [i for i, t in enumerate(texts) if len(t) < _PAGE_TEXT_SPARSE]
    if sparse and ocr:
        sparse = _filter_scan_pages(data, texts, sparse)
    candidates = sorted(set(sparse) | set(garbage))
    if candidates and ocr:
        texts = _ocr_fill_pages(data, texts, candidates, stats=stats,
                                on_ocr_progress=on_ocr_progress,
                                ocr_engine=ocr_engine, page_ocr=page_ocr)
    if stats is not None:
        stats["pages"] = len(texts)
        stats["ocrPages"] = len(candidates) if ocr else 0
        stats["ocrFilled"] = stats.get("ocrFilled", 0)
        if candidates and ocr:
            stats["ocrEngine"] = ocr_engine
        if ctrl_stripped:
            stats["ctrlStripped"] = ctrl_stripped
    _warn_watermark_scan(texts, ocr, stats)
    return "\n\n".join(t for t in texts if t)


# 一页文字层低于此字符数视为"疑似无文字层"(空页直接 OCR;非空短页
# 需结构证据——整页大图覆盖——才 OCR,见 _filter_scan_pages)
_PAGE_TEXT_SPARSE = 10


def _filter_scan_pages(data: bytes, texts: List[str],
                       sparse: List[int]) -> List[int]:
    """短文字层页结构复核:保留"空文字层"与"整页大图覆盖"两种候选。

    判据用结构证据(图片对象在版面上的覆盖矩形,对象遍历不解码位图),
    不再用"字符少"一概而论——章末空白页/整页图表短图注是合法稀疏页,
    误送 OCR 会把图内刻度、装饰线识别成文本,反噬原本正确的好图注。
    pypdfium2 不可用时维持全候选(与 OCR 依赖缺失的降级路径一致)。
    """
    try:
        import pypdfium2
        from pypdfium2.raw import FPDF_PAGEOBJ_IMAGE
    except ImportError:
        return sparse
    kept: List[int] = []
    pdf = pypdfium2.PdfDocument(io.BytesIO(data))
    try:
        for i in sparse:
            if not texts[i]:
                kept.append(i)  # 空文字层:无可失去,必候选
                continue
            page = pdf[i]
            try:
                w, h = page.get_size()
                covered = 0.0
                for obj in page.get_objects():
                    if obj.type == FPDF_PAGEOBJ_IMAGE:
                        l, b, r, t = obj.get_bounds()
                        covered = max(covered, (r - l) * (t - b))
                if w * h > 0 and covered / (w * h) >= _SCAN_IMAGE_COVERAGE:
                    kept.append(i)
            finally:
                page.close()
    finally:
        pdf.close()
    return kept


def _warn_watermark_scan(texts: List[str], ocr: bool,
                         stats: Dict[str, int]) -> None:
    """文档级盲区告警:整册扫描件若每页都带水印文字层,页级结构检查
    判不出"有字但全是水印"——按页字符中位数兜底提醒,只告警不阻断。"""
    if not ocr or stats is None or len(texts) < 4:
        return
    lens = sorted(len(t) for t in texts if t)
    if not lens or stats.get("ocrFilled"):
        return
    if lens[len(lens) // 2] < _PAGE_TEXT_SPARSE * 2:
        logger.warning(
            f"pdf 共 {len(texts)} 页,文字层页字符中位数仅 "
            f"{lens[len(lens) // 2]} 且 OCR 无补充——疑似整册扫描件"
            "(文字层只有水印),导入内容可能不完整")

# ── OCR 资源治理(单例 + 线程/并发封顶) ─────────────────────
# RapidOCR 默认 intra_op_num_threads=-1:det/cls/rec 三个 ONNX session 每个
# 都吃满全部核;而 TaskManager 允许不同库的导入并行(TASK_WORKERS,默认 2),
# 叠加起来就是线程超额订阅的 CPU 风暴 + 双份模型内存。三道闸:
# ① 进程级单例——模型只加载一份;② 线程封顶——牺牲单任务速度换整机稳定;
# ③ 推理锁——并行导入的 OCR 段逐页串行,任意时刻只有一张渲染位图在内存。
_OCR_INTRA_THREADS = 2
_OCR_RENDER_SCALE = 2.5  # ~180dpi,中文识别与速度的实测平衡点(上限)
_OCR_MAX_SIDE = 2000     # 对齐 RapidOCR Global.max_side_len 的内部降采样线
_OCR_LLM_CONCURRENCY = 3  # llm 引擎页级并发:网络调用无本地资源争抢
_OCR_ENGINE = None
_OCR_ENGINE_LOCK = threading.Lock()
_OCR_INFER_LOCK = threading.Lock()


def _ocr_engine():
    """进程级 RapidOCR 单例(懒加载,双检锁)。"""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        with _OCR_ENGINE_LOCK:
            if _OCR_ENGINE is None:
                from rapidocr_onnxruntime import RapidOCR
                _OCR_ENGINE = RapidOCR(intra_op_num_threads=_OCR_INTRA_THREADS)
    return _OCR_ENGINE


def _render_scale(page) -> float:
    """渲染倍率:2.5 封顶,且长边不超过 2000——超线的大图 RapidOCR 内部
    也会缩回 2000,先渲大图只是白占内存(A4@2.5 长边 2105px,缩掉 ~10%)。"""
    w, h = page.get_size()
    longest = max(w, h)
    return min(_OCR_RENDER_SCALE, _OCR_MAX_SIDE / longest) if longest > 0 else _OCR_RENDER_SCALE


def _ocr_fill_pages(data: bytes, texts: List[str], sparse_pages: List[int],
                    stats: Dict[str, int] = None, on_ocr_progress=None,
                    ocr_engine: str = "rapidocr", page_ocr=None) -> List[str]:
    """OCR 兜底:候选页(空文字层/整页大图的短文字层/文字层垃圾页)
    渲染成图 → 识别。

    ocr_engine: rapidocr(默认,本地 PP-OCR)/ llm(调用方注入的 page_ocr
    回调——任务层接多模态大模型;SDK 零领域知识,只认回调本身)。
    OCR 依赖未安装时保持原行为(返回原文字层),由任务层按"无有效文本"
    报错——不因可选依赖缺失而炸掉 md/txt/文字层 pdf 的解析。
    """
    import pypdfium2

    def _degrade(reason: str) -> List[str]:
        logger.warning(reason)
        if stats is not None:
            stats["ocrFilled"] = 0
        return texts

    if ocr_engine == "llm":
        if page_ocr is None:
            return _degrade("pdf OCR 引擎=llm 但未注入识别回调(page_ocr),跳过 OCR")

        def _recognize(img) -> str:
            return (page_ocr(img) or "").strip()
    else:
        try:
            import rapidocr_onnxruntime  # noqa: F401  仅探测;实例走 _ocr_engine() 单例
        except ImportError as e:
            return _degrade(f"pdf OCR 兜底不可用(未安装 rapidocr-onnxruntime): {e}")
        _ocr = _ocr_engine()

        def _recognize(img) -> str:
            result, _ = _ocr(img)
            return "\n".join(r[1] for r in (result or []))

    if on_ocr_progress:
        try:
            on_ocr_progress(0, len(sparse_pages))
        except Exception:  # noqa: BLE001  回调失败不影响识别
            pass
    pdf = pypdfium2.PdfDocument(io.BytesIO(data))
    filled = 0
    try:
        if ocr_engine == "llm":
            # llm 引擎:网络调用,无本地 CPU/内存争抢——页级并发(3 路)
            # 压缩 OCR 段耗时(14 页从 ~2 分钟压到 ~40 秒);渲染在主线程
            # 逐页做(pypdfium2 非线程安全),识别并发出去
            from concurrent.futures import ThreadPoolExecutor
            with ThreadPoolExecutor(max_workers=_OCR_LLM_CONCURRENCY) as pool:
                futures = []
                for i in sparse_pages:
                    page = pdf[i]
                    try:
                        img = page.render(scale=_render_scale(page)).to_pil()
                    except Exception as e:  # noqa: BLE001
                        logger.warning(f"pdf 第 {i + 1} 页渲染失败(跳过): {e}")
                        page.close()
                        continue
                    page.close()
                    futures.append((i, pool.submit(_recognize, img)))
                for n, (i, fut) in enumerate(futures):
                    try:
                        ocr_text = fut.result(timeout=300)
                    except Exception as e:  # noqa: BLE001
                        logger.warning(f"pdf 第 {i + 1} 页 OCR 失败(跳过): {e}")
                        continue
                    if len(ocr_text) > len(texts[i]):
                        texts[i] = ocr_text
                        filled += 1
                    if on_ocr_progress:
                        try:
                            on_ocr_progress(n + 1, len(sparse_pages))
                        except Exception:  # noqa: BLE001
                            pass
        else:
            # rapidocr 引擎:本地 CPU 推理,渲染+推理整体持锁——位图
            # ~10MB/页,并行导入时逐页排队,内存峰值只有一份
            for n, i in enumerate(sparse_pages):
                page = pdf[i]
                try:
                    with _OCR_INFER_LOCK:
                        img = page.render(scale=_render_scale(page)).to_pil()
                        ocr_text = _recognize(img)
                        del img
                except Exception as e:  # noqa: BLE001
                    logger.warning(f"pdf 第 {i + 1} 页 OCR 失败(跳过): {e}")
                    continue
                finally:
                    page.close()
                if len(ocr_text) > len(texts[i]):
                    texts[i] = ocr_text
                    filled += 1
                if on_ocr_progress:
                    try:
                        on_ocr_progress(n + 1, len(sparse_pages))
                    except Exception:  # noqa: BLE001
                        pass
    finally:
        pdf.close()
    if stats is not None:
        stats["ocrFilled"] = filled
    logger.info(f"pdf OCR 兜底({ocr_engine}): "
                f"{len(sparse_pages)} 页无文字层,识别成功 {filled} 页")
    return texts


def _parse_docx(data: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(data))
    parts: List[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if text:
            style = (para.style.name or "").lower()
            # 保留标题层级信息(切块的标题边界用)
            if "heading 1" in style:
                parts.append("# " + text)
            elif "heading 2" in style:
                parts.append("## " + text)
            elif "heading 3" in style:
                parts.append("### " + text)
            else:
                parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts)


# ── 切块:纯文本 → chunks ────────────────────────────────────

def _split_blocks(text: str) -> List[str]:
    """切成"结构块":标题行独立成块(保留 # 前缀),其余按空行分段。"""
    blocks: List[str] = []
    current: List[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if _is_heading_line(stripped):
            if current:
                blocks.append("\n".join(current).strip())
                current = []
            blocks.append(stripped)  # 标题独立成块(软边界)
        elif stripped == "":
            if current:
                blocks.append("\n".join(current).strip())
                current = []
        else:
            current.append(line)
    if current:
        blocks.append("\n".join(current).strip())
    return [b for b in blocks if b]


def _hard_split(block: str, max_chars: int) -> List[str]:
    """超长块按句读边界细切(句号/问叹号/分号/换行),仍超长则暴力截断。"""
    if len(block) <= max_chars:
        return [block]
    sentences = re.split(r"(?<=[。！？；;!?])\s*|\n", block)
    parts: List[str] = []
    buf = ""
    for sent in sentences:
        if not sent:
            continue
        while len(sent) > max_chars:  # 单句超长:暴力切
            if buf:
                parts.append(buf)
                buf = ""
            parts.append(sent[:max_chars])
            sent = sent[max_chars:]
        if len(buf) + len(sent) > max_chars and buf:
            parts.append(buf)
            buf = sent
        else:
            buf = buf + ("\n" if buf else "") + sent if buf else sent
    if buf:
        parts.append(buf)
    return parts


def sub_chunks_for_embedding(text: str, size: int = 500) -> List[str]:
    """向量化子切:按句读边界切成 ≤size 的段。

    抽取块可以是整章(structural_max 上限万级),而 embedding 输入有限
    (本地 bge-m3 实际按 4096 token 截断,超长部分对向量是隐形的;且单
    向量表达长文会语义稀释)。抽取按章、召回按段——子块 id 由调用方
    挂父块(f"{chunk_id}#i"),引用与去重按父块聚合。
    默认 500:Dify 对标口径——子段越小向量越聚焦、命中越准;连贯性由
    检索侧父子回带(命中子段回带父块上下文窗口)补足,不靠子段自身。
    """
    size = max(300, int(size or 500))
    return [p for p in _hard_split(text or "", size) if p.strip()]


def chunk_text(
    text: str,
    target_chars: int = 1200,
    overlap_chars: int = 100,
    max_chars: int = 3000,
    structural_max_chars: int = 10000,
) -> List[Dict[str, Any]]:
    """结构感知切块。Returns: [{seq, text, char_count}](seq 从 0 连续)。

    粒度策略(抽取与检索解耦的前提):
    - 结构块(标题开头:一章/一节)按 structural_max_chars 上限——能识别
      章节就不切,整章一块,超长章才在句读处兜底切;
    - 无结构纯文本维持 target_chars 软目标,max_chars 为其硬上限;
    - 向量化侧按需把结构块再子切(local_embeddings 实际仅 4096 token,
      长块尾部对向量是隐形的——抽取大块 + 向量小块,两头都最优)。
    """
    if not text or not text.strip():
        return []

    # 归一参数(防呆:非法配置不至于炸掉导入)
    target_chars = max(200, int(target_chars or 1200))
    max_chars = max(target_chars, int(max_chars or 3000))
    structural_max = max(max_chars, int(structural_max_chars or 10000))
    overlap_chars = max(0, min(int(overlap_chars or 0), target_chars // 2))

    chunks: List[str] = []
    buf = ""
    buf_structural = False   # 缓冲是否以标题开头(结构块:目录区/一章)
    buf_has_content = False  # 缓冲里是否混入过非标题正文(纯标题=目录行)
    for block in _split_blocks(text):
        is_head = _is_heading_line(block)
        # 标题 = 章节边界:上一缓冲含实质内容时结算,新章从标题重新起块。
        # 短于 _HEADING_FLUSH_MIN 的缓冲(孤立短行/单独"目录"两字)不值得
        # 独立成块,并入下一章。纯标题缓冲(目录页)不结算——目录行合并
        # 成尽量少的大块,而不是每行一个碎块(诛仙实测:逐行碎块让前 60+
        # 批 LLM 调用全为 0 实体)。
        if (is_head and buf and buf_has_content
                and len(buf) >= _HEADING_FLUSH_MIN):
            chunks.append(buf)
            buf = ""
            buf_structural = False
            buf_has_content = False
        if not is_head:
            # 正文到达而缓冲还是纯标题(目录区结束) → 目录自成一块结算,
            # 不与第一章正文粘连;短标题串(<40 字)并入正文。
            # 【坑位记录】标记必须在 pieces 循环前置位:放在循环后的话,
            # 章内第二个分段会把"标题+首段"误判成纯目录缓冲冲掉(实测踩过)
            if buf and not buf_has_content and len(buf) >= _HEADING_FLUSH_MIN:
                chunks.append(buf)
                buf = ""
                buf_structural = False
            buf_has_content = True
        for piece in _hard_split(block, max_chars):
            # 预切粒度 = max_chars:无结构路径的块硬上限不被突破;
            # 结构路径靠缓冲合并到 structural_max(碎件重组,不影响)
            # 结构块(标题开头)按 structural_max 上限:一章尽量一块——章内
            # 上下文完整,实体/关系不被章内切块打断;无结构纯文本维持
            # target_chars 软目标。超上限仍按句读细切兜底。
            ceiling = structural_max if (buf_structural or is_head) else target_chars
            if len(buf) + len(piece) + 1 > ceiling and buf:
                chunks.append(buf)
                buf = piece
                # 超限续段仍属同一章(结构上下文延续,标题在首块):
                # 保持结构标记,续段继续用 structural_max 上限,
                # 否则一章的后半会跌回 target 粒度被切碎
                buf_structural = buf_structural or is_head
            else:
                buf = buf + "\n" + piece if buf else piece
                buf_structural = buf_structural or is_head
    if buf:
        chunks.append(buf)

    # 相邻块重叠:后块带上前块尾部(跨块语义连续;标题块不加,避免污染)
    if overlap_chars > 0:
        overlapped: List[str] = []
        for i, c in enumerate(chunks):
            # 标题开头的块不加前块尾部。判定与切块软边界同源(_is_heading_line):
            # 曾只判 markdown #,中文章节标题块("第一章 xxx")被拼上上一章结尾,
            # 表现为"章标题在块中间"的假跨章
            if i == 0 or _is_heading_line(c.split("\n", 1)[0].strip()):
                overlapped.append(c)
                continue
            prev_tail = chunks[i - 1][-overlap_chars:]
            overlapped.append(prev_tail + "\n" + c)
        chunks = overlapped

    # 硬上限最终兜底(重叠可能超)
    final: List[str] = []
    for c in chunks:
        # 结构块允许到 structural_max,兜底切分用同一上限(无结构块
        # 本就 ≤ max_chars ≤ structural_max,不会被这里放大)
        final.extend(_hard_split(c, structural_max))

    return [
        {"seq": i, "text": t, "char_count": len(t)}
        for i, t in enumerate(final) if t.strip()
    ]
